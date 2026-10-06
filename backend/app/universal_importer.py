import io
import re
import logging
import xml.etree.ElementTree as ET
from typing import Optional

logger = logging.getLogger(__name__)

def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        return "\n\n".join(text_parts)
    except Exception as e:
        logger.error(f"Error extracting text from PDF: {str(e)}")
        raise ValueError(f"Failed to parse PDF document: {str(e)}")

def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        text_parts = []
        for para in doc.paragraphs:
            if para.text:
                text_parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text_parts.append(cell.text)
        return "\n".join(text_parts)
    except Exception as e:
        logger.error(f"Error extracting text from DOCX: {str(e)}")
        raise ValueError(f"Failed to parse DOCX document: {str(e)}")

def _fdx_paragraph_text(paragraph: ET.Element) -> str:
    """FDX splits a paragraph's text across multiple <Text> runs when
    formatting (bold/italic/underline) changes mid-line — concatenate all
    of them in document order to get the actual line."""
    return "".join(t.text or "" for t in paragraph.findall("Text")).strip()


def _fdx_to_pseudo_fountain(root: ET.Element) -> str:
    """
    Converts Final Draft's XML paragraph stream into the plain-text/Fountain
    shape app/parser.py already understands, so Stage 1 needs exactly one
    parser either way — this function's only job is format normalization,
    not screenplay structure extraction (that stays parser.py's job, kept
    deterministic per the spec).

    FDX <Paragraph Type="..."> values map onto parser.py's own detection
    rules: Scene Heading/Action/Character/Parenthetical/Dialogue/Transition.
    Character+Parenthetical+Dialogue paragraphs are siblings in the FDX
    stream (not nested), so they're re-grouped into one block here — that
    grouping is exactly what parser.py's block-splitting expects.
    """
    content = root.find("Content")
    if content is None:
        return ""

    blocks = []
    dialogue_block: list = []

    def flush_dialogue():
        if dialogue_block:
            blocks.append("\n".join(dialogue_block))
            dialogue_block.clear()

    for para in content.findall("Paragraph"):
        para_type = para.get("Type", "")
        text = _fdx_paragraph_text(para)
        if not text:
            continue

        if para_type == "Character":
            flush_dialogue()
            dialogue_block.append(text.upper())
        elif para_type in ("Parenthetical", "Dialogue"):
            if dialogue_block:
                line = f"({text})" if para_type == "Parenthetical" and not text.startswith("(") else text
                dialogue_block.append(line)
            else:
                # Orphan parenthetical/dialogue with no preceding Character
                # cue — fall back to an action line rather than drop it.
                flush_dialogue()
                blocks.append(text)
        elif para_type == "Scene Heading":
            flush_dialogue()
            # Force-heading prefix only if it doesn't already start with a
            # prefix parser.py recognizes — guards against non-standard
            # slugline wording still being detected as a heading.
            heading = text if re.match(r"^(INT|EXT|EST|I/E)\b", text, re.IGNORECASE) else f".{text}"
            blocks.append(heading)
        elif para_type == "Transition":
            flush_dialogue()
            blocks.append(text if text.upper().endswith("TO:") else f">{text}")
        else:
            # Action, Shot, General, or any other FDX paragraph type —
            # treat as action so nothing is silently dropped.
            flush_dialogue()
            blocks.append(text)

    flush_dialogue()
    return "\n\n".join(blocks)


def extract_text_from_fdx(file_bytes: bytes) -> str:
    try:
        root = ET.fromstring(file_bytes)
    except ET.ParseError as e:
        logger.error(f"Error parsing .fdx XML: {str(e)}")
        raise ValueError(f"Failed to parse Final Draft (.fdx) file — not valid XML: {str(e)}")

    text = _fdx_to_pseudo_fountain(root)
    if not text.strip():
        raise ValueError("Final Draft (.fdx) file parsed but contained no recognizable screenplay content.")
    return text


def extract_text(filename: str, file_bytes: bytes) -> str:
    fn_lower = filename.lower()
    if fn_lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif fn_lower.endswith(".docx") or fn_lower.endswith(".doc"):
        return extract_text_from_docx(file_bytes)
    elif fn_lower.endswith(".fdx"):
        return extract_text_from_fdx(file_bytes)
    else:
        # Fallback to text decoding
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                return file_bytes.decode("latin-1")
            except Exception as e:
                raise ValueError(f"Unable to decode text file: {str(e)}")
